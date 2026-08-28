from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


WRITE_HEADERS = {"X-Financial-System-Request": "1"}


def test_company_excel_template_preserves_formulas_and_removes_notes_below_row_five() -> None:
    template_path = Path(__file__).resolve().parents[2] / "新收费计划计算.xlsx"
    workbook = load_workbook(template_path, data_only=False)
    sheet = workbook["利润20%"]
    expected_main_formulas = {
        "K3": "=J3-I3+1",
        "O3": "=M3-N3",
        "P3": "=373455.51+94978.23",
        "Q3": "=P3-L3-O3",
        "R3": "=Q3/(L3+O3)",
        "T3": "=S3+O3",
        "U3": "=P3-T3",
        "V3": "=IF(U3>0,U3*0.2,0)",
        "W3": "=MAX(T3,P3)",
        "K4": "=J4-I4+1",
        "O4": "=M4-N4",
        "Q4": "=P4-L4-O4",
        "R4": "=Q4/(L4+O4)",
        "T4": "=S4+O4",
        "U4": "=P4-T4",
        "V4": "=IF(U4>0,U4*0.2,0)",
        "W4": "=MAX(T4,P4)",
    }
    assert sheet["C3"].value == "例子"
    assert sheet["C4"].value == "例子"
    assert {cell: sheet[cell].value for cell in expected_main_formulas} == expected_main_formulas
    personal_note_cells = {
        "D5", "F5", "H5", "M5", "N5", "O5", "P5", "Q5", "S5", "T5", "U5",
        "D6", "M6", "L7", "D8", "M9", "M10", "H12",
    }
    assert all(sheet[address].value is None for address in personal_note_cells)
    chinese_cells = {
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and re.search(r"[\u4e00-\u9fff]", cell.value)
    }
    assert chinese_cells == {"C3", "C4"}
    defer = workbook["Defer 延付利息（待确认）"]
    defer_formulas = [
        cell.value
        for row in defer.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert len(defer_formulas) == 36
    with ZipFile(template_path) as package:
        assert {
            "docProps/app.xml",
            "docProps/core.xml",
            "docProps/custom.xml",
            "xl/calcChain.xml",
            "xl/theme/theme1.xml",
        }.issubset(package.namelist())


def _master(client: TestClient, suffix: str, *, accounts: int = 1) -> dict:
    company = client.post("/api/companies", json={"name": f"Account HWM {suffix}", "code": f"AH{suffix}"}).json()
    fc = client.post("/api/fcs", json={"company_id": company["id"], "name": f"FC {suffix}", "code": f"F{suffix}"}).json()
    platform = client.post("/api/platforms", json={"name": f"Platform {suffix}", "code": f"P{suffix}"}).json()
    plan = client.post(
        "/api/fee-plans",
        json={"company_id": company["id"], "name": "利润20%", "code": f"PL{suffix}", "fee_rate_percent": "20"},
    ).json()
    customer = client.post(
        "/api/clients",
        json={
            "company_id": company["id"],
            "fc_id": fc["id"],
            "name": f"Client {suffix}",
            "management_start_date": "2026-01-01",
            "status": "ACTIVE",
        },
    ).json()
    created_accounts = []
    for number in range(1, accounts + 1):
        created_accounts.append(
            client.post(
                "/api/accounts",
                json={
                    "client_id": customer["id"],
                    "platform_id": platform["id"],
                    "fee_plan_id": plan["id"],
                    "account_number": f"{suffix}-{number}",
                    "start_date": "2026-01-01",
                    "status": "ACTIVE",
                },
            ).json()
        )
    return {"company": company, "fc": fc, "platform": platform, "plan": plan, "client": customer, "accounts": created_accounts}


def _attach(client: TestClient, entity_type: str, entity_id: int) -> None:
    response = client.post(
        "/api/attachments",
        data={"entity_type": entity_type, "entity_id": str(entity_id)},
        files={"file": (f"proof-{entity_type}-{entity_id}.pdf", b"%PDF-1.4\nproof\n%%EOF", "application/pdf")},
    )
    assert response.status_code == 201, response.text


def _snapshot(
    client: TestClient,
    account_id: int,
    as_of_date: str,
    balance: str,
    *,
    closing: bool,
    evidence: bool,
) -> dict:
    response = client.post(
        "/api/balance-snapshots",
        json={
            "account_id": account_id,
            "as_of_date": as_of_date,
            "total_balance": balance,
            "eligible_for_closing": closing,
        },
    )
    assert response.status_code == 201, response.text
    item = response.json()
    if evidence:
        _attach(client, "SNAPSHOT", item["id"])
    return item


def _calculate(client: TestClient, data: dict, lines: list[dict], *, quarter: int = 1) -> dict:
    response = client.post(
        "/api/settlements/calculate",
        json={
            "client_id": data["client"]["id"],
            "platform_id": data["platform"]["id"],
            "fee_plan_id": data["plan"]["id"],
            "year": 2026,
            "quarter": quarter,
            "account_lines": lines,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_each_account_fee_is_calculated_before_group_total() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        data = _master(client, "OFFSET", accounts=2)
        lines = []
        for account, closing_balance in zip(data["accounts"], ["1100.00", "800.00"], strict=True):
            beginning = _snapshot(client, account["id"], "2026-01-01", "1000.00", closing=False, evidence=True)
            closing = _snapshot(client, account["id"], "2026-03-31", closing_balance, closing=True, evidence=True)
            lines.append(
                {
                    "account_id": account["id"],
                    "beginning_snapshot_id": beginning["id"],
                    "closing_snapshot_id": closing["id"],
                    "original_hwm": "1000.00",
                }
            )
        settlement = _calculate(client, data, lines)
        assert settlement["calculation_mode"] == "ACCOUNT_HWM"
        assert settlement["formula_version"] == "HWM-2.0-ACCOUNT"
        assert [line["service_fee"] for line in settlement["account_lines"]] == ["20.00", "0.00"]
        assert settlement["watermark_difference"] == "-100.00"
        assert settlement["chargeable_above_hwm"] == "100.00"
        assert settlement["service_fee"] == "20.00"
        finalized = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert finalized.status_code == 200, finalized.text
        excel = client.post(f"/api/exports/excel?settlement_ids={settlement['id']}")
        assert excel.status_code == 200, excel.text
        sheet = load_workbook(BytesIO(excel.content), data_only=False)["利润20%"]
        assert [sheet["H3"].value, sheet["H4"].value] == ["OFFSET-1", "OFFSET-2"]
        assert [sheet["V3"].value, sheet["V4"].value] == [20, 0]
        pdf = client.post(f"/api/exports/pdf?settlement_id={settlement['id']}&language=zh")
        assert pdf.status_code == 200, pdf.text


def test_internal_finance_excel_batches_selected_finalized_settlements() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        settlement_ids = []
        for suffix, closing_balance in [("BATCHA", "1100.00"), ("BATCHB", "1200.00")]:
            data = _master(client, suffix)
            account = data["accounts"][0]
            beginning = _snapshot(
                client, account["id"], "2026-01-01", "1000.00", closing=False, evidence=True
            )
            closing = _snapshot(
                client, account["id"], "2026-03-31", closing_balance, closing=True, evidence=True
            )
            settlement = _calculate(
                client,
                data,
                [{
                    "account_id": account["id"],
                    "beginning_snapshot_id": beginning["id"],
                    "closing_snapshot_id": closing["id"],
                    "original_hwm": "1000.00",
                }],
            )
            finalized = client.post(f"/api/settlements/{settlement['id']}/finalize")
            assert finalized.status_code == 200, finalized.text
            settlement_ids.append(settlement["id"])

        response = client.post(
            f"/api/exports/excel?settlement_ids={settlement_ids[1]},{settlement_ids[0]}"
        )
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        assert workbook.sheetnames == ["利润20%", "Defer 延付利息（待确认）"]
        sheet = workbook["利润20%"]
        assert [sheet["E3"].value, sheet["E4"].value] == ["Client BATCHA", "Client BATCHB"]
        assert [sheet["H3"].value, sheet["H4"].value] == ["BATCHA-1", "BATCHB-1"]
        assert [sheet["V3"].value, sheet["V4"].value] == [20, 40]
        assert sheet["K3"].value == "=J3-I3+1"
        assert sheet["O3"].value == "=M3-N3"
        assert sheet["Q3"].value == "=P3-L3-O3"
        assert sheet["R3"].value == '=IFERROR(Q3/(L3+O3),"")'
        assert sheet["T3"].value == "=S3+O3"
        assert sheet["U3"].value == "=P3-T3"
        assert sheet["W3"].value == "=MAX(T3,P3)"
        assert all(sheet.cell(3, column).alignment.wrap_text for column in range(2, 9))
        assert workbook["Defer 延付利息（待确认）"]["X3"].value is not None


def test_each_account_uses_its_own_starting_and_closing_dates() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        data = _master(client, "ACCTPERIOD", accounts=2)
        first, second = data["accounts"]
        updated = client.patch(
            f"/api/accounts/{second['id']}", json={"end_date": "2026-02-28"}
        )
        assert updated.status_code == 200, updated.text

        first_beginning = _snapshot(client, first["id"], "2026-01-01", "1000.00", closing=False, evidence=True)
        first_closing = _snapshot(client, first["id"], "2026-03-31", "1100.00", closing=True, evidence=True)
        second_beginning = _snapshot(client, second["id"], "2026-02-01", "2000.00", closing=False, evidence=True)
        second_closing = _snapshot(client, second["id"], "2026-02-28", "2200.00", closing=True, evidence=True)
        before_second_account_period = client.post(
            "/api/transactions",
            json={
                "account_id": second["id"],
                "transaction_date": "2026-01-15",
                "transaction_type": "CONTRIBUTION",
                "amount": "500.00",
            },
        )
        assert before_second_account_period.status_code == 201, before_second_account_period.text

        settlement = _calculate(
            client,
            data,
            [
                {
                    "account_id": first["id"],
                    "start_date": "2026-01-01",
                    "closing_date": "2026-03-31",
                    "beginning_snapshot_id": first_beginning["id"],
                    "closing_snapshot_id": first_closing["id"],
                    "original_hwm": "1000.00",
                },
                {
                    "account_id": second["id"],
                    "start_date": "2026-02-01",
                    "closing_date": "2026-02-28",
                    "beginning_snapshot_id": second_beginning["id"],
                    "closing_snapshot_id": second_closing["id"],
                    "original_hwm": "2000.00",
                },
            ],
        )

        assert settlement["start_date"] == "2026-01-01"
        assert settlement["closing_date"] == "2026-03-31"
        assert [
            (line["start_date"], line["closing_date"], line["days"])
            for line in settlement["account_lines"]
        ] == [
            ("2026-01-01", "2026-03-31", 90),
            ("2026-02-01", "2026-02-28", 28),
        ]
        assert settlement["account_lines"][1]["contribution"] == "0.00"
        finalized = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert finalized.status_code == 200, finalized.text

        after_second_account_period = client.post(
            "/api/transactions",
            json={
                "account_id": second["id"],
                "transaction_date": "2026-03-15",
                "transaction_type": "CONTRIBUTION",
                "amount": "100.00",
            },
        )
        assert after_second_account_period.status_code == 201, after_second_account_period.text

        excel = client.post(f"/api/exports/excel?settlement_ids={settlement['id']}")
        assert excel.status_code == 200, excel.text
        sheet = load_workbook(BytesIO(excel.content), data_only=False)["利润20%"]
        assert sheet["I3"].value.strftime("%Y-%m-%d") == "2026-01-01"
        assert sheet["J3"].value.strftime("%Y-%m-%d") == "2026-03-31"
        assert sheet["I4"].value.strftime("%Y-%m-%d") == "2026-02-01"
        assert sheet["J4"].value.strftime("%Y-%m-%d") == "2026-02-28"


def test_draft_allows_missing_snapshot_evidence_but_finalize_blocks_it() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        data = _master(client, "EVIDENCE")
        account = data["accounts"][0]
        beginning = _snapshot(client, account["id"], "2026-01-01", "1000.00", closing=False, evidence=False)
        closing = _snapshot(client, account["id"], "2026-03-31", "1100.00", closing=True, evidence=False)
        settlement = _calculate(
            client,
            data,
            [{
                "account_id": account["id"],
                "beginning_snapshot_id": beginning["id"],
                "closing_snapshot_id": closing["id"],
                "original_hwm": "1000.00",
            }],
        )
        blocked = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert blocked.status_code == 400
        assert "缺少凭证" in blocked.json()["detail"]
        _attach(client, "SNAPSHOT", beginning["id"])
        _attach(client, "SNAPSHOT", closing["id"])
        finalized = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert finalized.status_code == 200, finalized.text


def test_one_zero_denominator_account_blocks_group_finalize() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        data = _master(client, "ZERODENOM", accounts=2)
        lines = []
        for account, balance in zip(data["accounts"], ["1000.00", "0.00"], strict=True):
            beginning = _snapshot(client, account["id"], "2026-01-01", balance, closing=False, evidence=True)
            closing = _snapshot(client, account["id"], "2026-03-31", balance, closing=True, evidence=True)
            lines.append({
                "account_id": account["id"],
                "beginning_snapshot_id": beginning["id"],
                "closing_snapshot_id": closing["id"],
                "original_hwm": balance,
            })

        settlement = _calculate(client, data, lines)
        assert settlement["period_rate"] == 0
        assert settlement["account_lines"][1]["period_rate"] is None
        blocked = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert blocked.status_code == 400
        assert data["accounts"][1]["account_number"] in blocked.json()["detail"]


def test_transaction_requires_its_own_evidence_before_finalize() -> None:
    with TestClient(app, headers=WRITE_HEADERS) as client:
        data = _master(client, "TXPROOF")
        account = data["accounts"][0]
        beginning = _snapshot(client, account["id"], "2026-01-01", "1000.00", closing=False, evidence=True)
        closing = _snapshot(client, account["id"], "2026-03-31", "1210.00", closing=True, evidence=True)
        transaction = client.post(
            "/api/transactions",
            json={
                "account_id": account["id"],
                "transaction_date": "2026-02-15",
                "transaction_type": "CONTRIBUTION",
                "amount": "100.00",
            },
        ).json()
        settlement = _calculate(
            client,
            data,
            [{
                "account_id": account["id"],
                "beginning_snapshot_id": beginning["id"],
                "closing_snapshot_id": closing["id"],
                "original_hwm": "1000.00",
            }],
        )
        assert settlement["account_lines"][0]["contribution"] == "100.00"
        blocked = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert blocked.status_code == 400
        assert "CONTRIBUTION" in blocked.json()["detail"]
        _attach(client, "TRANSACTION", transaction["id"])
        finalized = client.post(f"/api/settlements/{settlement['id']}/finalize")
        assert finalized.status_code == 200, finalized.text
