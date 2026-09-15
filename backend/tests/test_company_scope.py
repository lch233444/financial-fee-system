from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from pypdf import PdfReader
from openpyxl import load_workbook

from app.main import app
from test_invoice_aggregation import WRITE_HEADERS, _finalized_settlement, _group
from test_payment_corrections import _unclaimed_proof


def create(api, path, data):
    result = api.post(path, json=data)
    assert result.status_code == 201, result.text
    return result.json()


def test_company_is_not_needed_until_bill_and_chosen_payee_drives_archive_and_cash():
    tag = uuid4().hex[:12]
    with TestClient(app, headers=WRITE_HEADERS) as api:
        fc = create(api, '/api/fcs', {'name': 'Independent FC', 'code': 'F' + tag})
        plan = create(api, '/api/fee-plans', {'name': 'Independent Plan', 'code': 'P' + tag})
        client = create(api, '/api/clients', {'name': 'Independent ' + tag, 'fc_id': fc['id'],
                        'management_start_date': '2026-01-01', 'status': 'ACTIVE'})
        assert fc['company_id'] is plan['company_id'] is client['company_id'] is None
        platform = create(api, '/api/platforms', {'name': 'Independent ' + tag, 'code': 'PL' + tag})
        account = create(api, '/api/accounts', {'client_id': client['id'], 'platform_id': platform['id'],
                         'fee_plan_id': plan['id'], 'account_number': tag, 'start_date': '2026-01-01', 'status': 'ACTIVE'})
        group = {'client': client, 'fc': fc, 'plan': plan, 'platforms': [platform], 'accounts': [account]}
        settlement = _finalized_settlement(api, group, 0, year=2026, quarter=1)
        assert settlement['company_id'] is None and settlement['service_fee'] == '20.00'
        before_excel = api.post(f"/api/exports/excel?settlement_ids={settlement['id']}")
        assert before_excel.status_code == 200
        workbook = load_workbook(BytesIO(before_excel.content))
        assert workbook['收费计算']['B3'].value is None
        workbook.close()
        body = {'client_id': client['id'], 'year': 2026, 'quarter': 1}
        missing = api.post('/api/invoices', json=body)
        assert missing.status_code == 400 and '收款公司' in missing.text
        assert api.post('/api/invoices', json={**body, 'payee_company_id': 2147483647}).status_code == 404
        company = create(api, '/api/companies', {'name': 'Selected Payee ' + tag, 'code': tag,
                         'bank_information': 'Chosen bank 987654', 'payment_terms_days': 30})
        draft = create(api, '/api/invoices', {**body, 'payee_company_id': company['id']})
        issued = api.post(f"/api/invoices/{draft['id']}/issue", json={'issue_date': '2026-04-05'})
        assert issued.status_code == 200, issued.text
        bill = issued.json()
        assert bill['payee_company_id'] == company['id'] and bill['due_date'] == '2026-05-05'
        assert bill['invoice_number'].isdigit() and len(bill['invoice_number']) == 9 and bill['amount'] == '20.00'
        after_excel = api.post(f"/api/exports/excel?settlement_ids={settlement['id']}")
        assert after_excel.status_code == 200
        workbook = load_workbook(BytesIO(after_excel.content))
        assert workbook['收费计算']['B3'].value == company['name']
        assert workbook['收费计算']['D3'].value == bill['invoice_number']
        workbook.close()
        pdf = api.post(f"/api/invoices/{bill['id']}/pdf?language=en").content
        content = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
        assert company['name'] in content and '987654' in content
        proof = _unclaimed_proof(api, 'PAYMENT', tag)
        create(api, f"/api/invoices/{bill['id']}/payments", {'payment_date': '2026-04-10', 'amount': '20.00',
                       'method': 'BANK_TRANSFER', 'proof_attachment_id': proof})
        assert api.get(f"/api/invoices/{bill['id']}").json()['payment_status'] == 'PAID'
        after = next(row for row in api.get('/api/settlements').json() if row['id'] == settlement['id'])
        assert after == settlement


def test_legacy_company_metadata_does_not_restrict_fc_plan_or_payee():
    with TestClient(app, headers=WRITE_HEADERS) as api:
        a = _group(api, uuid4().hex[:10], platform_count=1)
        b = _group(api, uuid4().hex[:10], platform_count=1)
        # Old ownership differs in all three records, without changing FC identity.
        patched = api.patch(f"/api/clients/{a['client']['id']}", json={'company_id': b['company']['id']})
        assert patched.status_code == 200, patched.text
        patched = api.patch(f"/api/accounts/{a['accounts'][0]['id']}", json={'fee_plan_id': b['plan']['id']})
        assert patched.status_code == 200, patched.text
        a['plan'] = b['plan']
        settlement = _finalized_settlement(api, a, 0, year=2026, quarter=1)
        assert settlement['company_id'] is None and settlement['fc_id'] == a['fc']['id']
        draft = create(api, '/api/invoices', {'client_id': a['client']['id'], 'year': 2026, 'quarter': 1,
                       'payee_company_id': b['company']['id']})
        result = api.post(f"/api/invoices/{draft['id']}/issue", json={'issue_date': '2026-04-05'})
        assert result.status_code == 200, result.text
        assert result.json()['company_name'] == b['company']['name']


def test_new_codes_are_globally_unique_including_concurrent_creation():
    with TestClient(app, headers=WRITE_HEADERS) as api:
        tag = uuid4().hex[:12]
        company = create(api, '/api/companies', {'name': tag, 'code': tag})
        for path in ('/api/fcs', '/api/fee-plans'):
            body = {'name': 'Concurrent ' + tag, 'code': tag}
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = [pool.submit(api.post, path, json=values) for values in
                           (body, {**body, 'company_id': company['id']})]
                assert sorted(result.result().status_code for result in results) == [201, 409]
            duplicate = api.post(path, json={**body, 'code': tag.lower()})
            assert duplicate.status_code == 409, duplicate.text


def test_customer_still_requires_fc_and_start_date_for_activation():
    with TestClient(app, headers=WRITE_HEADERS) as api:
        assert api.post('/api/clients', json={'name': uuid4().hex, 'status': 'ACTIVE',
                        'management_start_date': '2026-01-01'}).status_code == 400
