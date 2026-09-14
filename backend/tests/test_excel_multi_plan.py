"""Exercise the real export endpoint using isolated synthetic settlements."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.database import SessionLocal
from app.main import app
from app.models import ExportRecord, FeePlan
from test_invoice_aggregation import WRITE_HEADERS, _group, _finalized_settlement, _draft
from test_invoice_multi_plan import _second_plan, _issue_draft


def _fee(value):
    return Decimal(str(value))


def test_mixed_plan_export_matches_each_locked_account_and_issued_invoice():
    with TestClient(app, headers=WRITE_HEADERS) as client:
        first = _group(client, "XLSMIX", platform_count=2)
        second = _second_plan(client, first, "XLSMIX", rate="30.00")
        other = _group(client, "XLSOTHER", platform_count=1)
        results = [
            _finalized_settlement(client, first, 0, year=2026, quarter=1, closing="1100.03"),
            _finalized_settlement(client, first, 1, year=2026, quarter=1, closing="700.00"),
            _finalized_settlement(client, second, 0, year=2026, quarter=1, closing="1300.05"),
        ]
        _finalized_settlement(client, other, 0, year=2026, quarter=1, closing="2000.00")
        issued = _issue_draft(client, _draft(client, first, year=2026, quarter=1))
        assert issued["amount"] == "110.03"
        # Deliberately change only this synthetic master record to distinguish
        # a live plan rate from the finalized account's historical rate.
        with SessionLocal() as db:
            db.get(FeePlan, second["plan"]["id"]).fee_rate_bps = 2500
            db.commit()
        ids = ",".join(str(result["id"]) for result in reversed(results))
        response = client.post(f"/api/exports/excel?settlement_ids={ids}")
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        assert workbook.sheetnames == ["收费计算", "Defer 延付利息（待确认）"]
        sheet = workbook["收费计算"]
        assert "20%" not in sheet["V2"].value
        assert sheet["Z1"].value == "Fee Plan / 收费计划"
        assert sheet["AA1"].value == "Fee Rate / 收费费率"
        for row, result in enumerate(results, start=3):
            line = result["account_lines"][0]
            plan = second["plan"] if result["fee_plan_id"] == second["plan"]["id"] else first["plan"]
            assert sheet.cell(row, 2).value == first["company"]["name"]
            assert sheet.cell(row, 3).value == "2026 Q1"
            assert sheet.cell(row, 4).value == issued["invoice_number"]
            assert sheet.cell(row, 5).value == first["client"]["name"]
            assert sheet.cell(row, 6).value == first["fc"]["name"]
            assert sheet.cell(row, 7).value == result["platform_name"]
            assert sheet.cell(row, 8).value == line["account_number"]
            assert sheet.cell(row, 9).value.date().isoformat() == line["start_date"]
            assert sheet.cell(row, 10).value.date().isoformat() == line["closing_date"]
            for col, key in [(12, "beginning"), (13, "contribution"), (14, "withdrawal"), (16, "closing"), (19, "original_hwm"), (22, "service_fee")]:
                assert _fee(sheet.cell(row, col).value) == _fee(line[key]), (row, key)
            assert sheet.cell(row, 24).value.date().isoformat() == issued["issue_date"]
            assert sheet.cell(row, 25).value.date().isoformat() == issued["due_date"]
            assert sheet.cell(row, 26).value == f"{plan['name']} ({plan['code']})"
            assert _fee(sheet.cell(row, 27).value) == _fee(result["fee_rate"])
            assert sheet.cell(row, 27).number_format == "0.00%"
            assert sheet.cell(row, 26).alignment.wrap_text
        assert [sheet.cell(row, 22).value for row in range(3, 6)] == [20.01, 0, 90.02]
        assert sheet["AA5"].value == 0.3
        assert sheet.auto_filter.ref == "A1:AA5"
        assert all(cell.value is None for cell in sheet._cells.values() if cell.row > 5)
        template = load_workbook(Path(__file__).resolve().parents[2] / "新收费计划计算纯净版模板.xlsx")
        defer = "Defer 延付利息（待确认）"
        expected = {key: (cell.value, cell.number_format) for key, cell in template[defer]._cells.items()}
        actual = {key: (cell.value, cell.number_format) for key, cell in workbook[defer]._cells.items()}
        assert actual == expected


def test_export_plan_filter_rejects_hidden_other_plan_and_exports_exact_selection():
    with TestClient(app, headers=WRITE_HEADERS) as client:
        first = _group(client, "XLSFILTER", platform_count=1)
        second = _second_plan(client, first, "XLSFILTER", rate="30.00")
        s1 = _finalized_settlement(client, first, 0, year=2026, quarter=1)
        s2 = _finalized_settlement(client, second, 0, year=2026, quarter=1)
        plan_id = second["plan"]["id"]
        with SessionLocal() as db:
            before = db.scalar(select(func.count()).select_from(ExportRecord))
        rejected = client.post(f"/api/exports/excel?settlement_ids={s1['id']},{s2['id']}&fee_plan_id={plan_id}")
        assert rejected.status_code == 400
        assert "其他收费计划" in rejected.json()["detail"]
        with SessionLocal() as db:
            assert db.scalar(select(func.count()).select_from(ExportRecord)) == before
        response = client.post(f"/api/exports/excel?settlement_ids={s2['id']},{s2['id']}&fee_plan_id={plan_id}")
        assert response.status_code == 200, response.text
        sheet = load_workbook(BytesIO(response.content))["收费计算"]
        assert sheet["H3"].value == second["accounts"][0]["account_number"]
        assert sheet["V3"].value == 30
        assert sheet["AA3"].value == 0.3
        assert sheet["H4"].value is None
        assert sheet["B3"].value is None and sheet["D3"].value is None
        assert sheet["X3"].value is None and sheet["Y3"].value is None
